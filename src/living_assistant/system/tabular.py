"""Read column templates and append rows to CSV / XLSX files.

Rows are dicts keyed by column name. When the target file already exists its own
header row decides the column order and which columns are filled (matched
case- and punctuation-insensitively), so a user's template keeps its layout,
extra columns and, for XLSX, its formatting and other sheets.

Scraped text is untrusted: a cell that starts with = + - @ would run as a
formula when the file is opened in a spreadsheet, so such values are stored as
plain text.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

SUPPORTED_SUFFIXES = {".csv", ".xlsx"}
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_NUMBER_RE = re.compile(r"-?\d+(\.\d+)?")


def normalize_column(name: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name or "").lower())


def check_suffix(path: Path) -> None:
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported table format {path.suffix or '(none)'}; use .csv or .xlsx.")


def read_headers(path: Path) -> list[str]:
    check_suffix(path)
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            row = next(csv.reader(fh), [])
    else:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        try:
            first = next(wb.active.iter_rows(min_row=1, max_row=1, values_only=True), ())
            row = list(first)
        finally:
            wb.close()
    return [str(h).strip() for h in row if h is not None and str(h).strip()]


def read_rows(path: Path) -> tuple[list[str], list[dict]]:
    """Return (headers, rows) of an existing table file."""
    check_suffix(path)
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.reader(fh)
            headers = [h.strip() for h in next(reader, [])]
            rows = [dict(zip(headers, r)) for r in reader if any(str(v).strip() for v in r)]
        return headers, rows
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        it = wb.active.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it, ())]
        rows = []
        for r in it:
            if r and any(v not in (None, "") for v in r):
                rows.append({h: ("" if v is None else str(v)) for h, v in zip(headers, r) if h})
    finally:
        wb.close()
    return headers, rows


def _text_safe(value: str) -> str:
    """CSV cells: prefix a quote so spreadsheets never evaluate scraped text."""
    if value.startswith(_FORMULA_START) and not _NUMBER_RE.fullmatch(value):
        return "'" + value
    return value


def _row_key(row: dict, headers: list[str]) -> tuple:
    return tuple(str(row.get(h, "")).strip().lower() for h in headers)


def _map_row(row: dict, headers: list[str]) -> dict:
    by_norm = {normalize_column(k): v for k, v in row.items()}
    return {h: str(by_norm.get(normalize_column(h), "") or "") for h in headers}


def write_rows(path: Path, columns: list[str], rows: list[dict], append: bool = True) -> dict:
    """Write ``rows`` to ``path``; append (skipping duplicates) when it exists and ``append``."""
    check_suffix(path)
    existing = path.exists() and path.stat().st_size > 0 and append
    if existing:
        headers, old_rows = read_rows(path)
        headers = [h for h in headers if h] or list(columns)
    else:
        headers, old_rows = list(columns), []

    seen = {_row_key(_map_row(r, headers), headers) for r in old_rows}
    fresh, duplicates = [], 0
    for row in rows:
        mapped = _map_row(row, headers)
        if not any(v.strip() for v in mapped.values()):
            continue
        key = _row_key(mapped, headers)
        if key in seen:
            duplicates += 1
            continue
        seen.add(key)
        fresh.append(mapped)

    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        _write_csv(path, headers, fresh, existing)
    else:
        _write_xlsx(path, headers, fresh, existing)
    return {
        "path": str(path),
        "columns": headers,
        "added": len(fresh),
        "skipped_duplicates": duplicates,
        "total_rows": len(old_rows) + len(fresh),
    }


def _write_csv(path: Path, headers: list[str], rows: list[dict], existing: bool) -> None:
    if existing:
        with path.open("rb") as fh:
            fh.seek(0, 2)
            needs_newline = False
            if fh.tell() > 0:
                fh.seek(-1, 2)
                needs_newline = fh.read(1) not in (b"\n", b"\r")
        with path.open("a", newline="", encoding="utf-8") as fh:
            if needs_newline:
                fh.write("\r\n")
            writer = csv.writer(fh)
            for row in rows:
                writer.writerow([_text_safe(row[h]) for h in headers])
        return
    # utf-8-sig so Excel opens non-ASCII text (prices in £/€, names) correctly.
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([_text_safe(row[h]) for h in headers])


def _set_cell(cell, value: str) -> None:
    if _NUMBER_RE.fullmatch(value) and len(value) < 16 and not (len(value) > 1 and value.lstrip("-").startswith("0") and "." not in value):
        cell.value = float(value) if "." in value else int(value)
        return
    cell.value = value
    cell.data_type = "s"  # openpyxl would otherwise store "=..." as a live formula
    if value.startswith(("http://", "https://")) and len(value) < 2000:
        cell.hyperlink = value
        cell.style = "Hyperlink"


def _write_xlsx(path: Path, headers: list[str], rows: list[dict], existing: bool) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if existing:
        wb = load_workbook(path)
        ws = wb.active
        header_cells = {str(c.value).strip(): c.column for c in ws[1] if c.value not in (None, "")}
        # Templates often pre-format hundreds of empty rows; append after the last one holding data.
        last = 1
        for r in range(ws.max_row, 1, -1):
            if any(ws.cell(r, c).value not in (None, "") for c in header_cells.values()):
                last = r
                break
        for i, row in enumerate(rows, start=last + 1):
            for h in headers:
                if h in header_cells and row[h]:
                    _set_cell(ws.cell(i, header_cells[h]), row[h])
        wb.save(path)
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(headers)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="3B4FB8")
    for cell in ws[1]:
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(vertical="center")
    for i, row in enumerate(rows, start=2):
        for j, h in enumerate(headers, start=1):
            if row[h]:
                _set_cell(ws.cell(i, j), row[h])
    for j, h in enumerate(headers, start=1):
        width = max([len(h)] + [len(r[h]) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(60, width + 2))
    ws.freeze_panes = "A2"
    if headers:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    wb.save(path)
